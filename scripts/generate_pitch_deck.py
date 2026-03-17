from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
LOGO_PATH = ROOT / "static" / "netcopilot-logo.png"
ASSETS_DIR = ROOT / "sales" / "assets"
OUTPUT_PATH = ROOT / "sales" / "NetCopilot_Pitch_Deck.pptx"
PRESENTATION_NOTE_URL = "https://gentle-harmony.lovable.app"


@dataclass(frozen=True)
class Palette:
    ink: RGBColor = RGBColor(11, 27, 51)  # deep navy
    slate: RGBColor = RGBColor(51, 65, 85)
    muted: RGBColor = RGBColor(100, 116, 139)
    bg: RGBColor = RGBColor(245, 247, 255)
    panel: RGBColor = RGBColor(255, 255, 255)
    blue: RGBColor = RGBColor(26, 142, 242)
    indigo: RGBColor = RGBColor(110, 86, 207)
    border: RGBColor = RGBColor(210, 225, 255)


P = Palette()


def _ensure_assets() -> Path:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    return ASSETS_DIR


def _make_noc_engineer_icon(path: Path) -> None:
    """Creates a simple 'NOC engineer at console with headset' icon as a PNG."""
    size = (900, 900)
    img = Image.new("RGBA", size, (0, 0, 0, 0))
    d = ImageDraw.Draw(img)

    # Background blob (soft gradient-like via layered circles)
    for r, alpha in [(420, 35), (360, 45), (300, 55)]:
        d.ellipse(
            (size[0] // 2 - r, size[1] // 2 - r, size[0] // 2 + r, size[1] // 2 + r),
            fill=(26, 142, 242, alpha),
        )
    for r, alpha in [(320, 25), (260, 35), (220, 45)]:
        d.ellipse(
            (size[0] // 2 - r + 140, size[1] // 2 - r - 120, size[0] // 2 + r + 140, size[1] // 2 + r - 120),
            fill=(110, 86, 207, alpha),
        )

    # Head
    head_center = (340, 360)
    head_r = 110
    d.ellipse(
        (head_center[0] - head_r, head_center[1] - head_r, head_center[0] + head_r, head_center[1] + head_r),
        fill=(255, 255, 255, 255),
        outline=(210, 225, 255, 255),
        width=6,
    )

    # Headset band
    d.arc(
        (head_center[0] - 140, head_center[1] - 150, head_center[0] + 140, head_center[1] + 110),
        start=200,
        end=-20,
        fill=(26, 142, 242, 255),
        width=14,
    )
    # Earpads
    d.rounded_rectangle((head_center[0] - 155, head_center[1] - 40, head_center[0] - 105, head_center[1] + 40), radius=18, fill=(230, 240, 255, 255))
    d.rounded_rectangle((head_center[0] + 105, head_center[1] - 40, head_center[0] + 155, head_center[1] + 40), radius=18, fill=(230, 240, 255, 255))

    # Mic boom
    d.line((head_center[0] + 120, head_center[1] + 30, head_center[0] + 210, head_center[1] + 90), fill=(26, 142, 242, 255), width=10)
    d.ellipse((head_center[0] + 202, head_center[1] + 82, head_center[0] + 230, head_center[1] + 110), fill=(110, 86, 207, 255))

    # Body / hoodie
    d.rounded_rectangle((190, 470, 490, 760), radius=90, fill=(255, 255, 255, 255), outline=(210, 225, 255, 255), width=6)
    d.polygon([(340, 520), (260, 610), (420, 610)], fill=(240, 246, 255, 255))

    # Console monitor
    d.rounded_rectangle((520, 340, 840, 560), radius=30, fill=(11, 27, 51, 255))
    d.rounded_rectangle((545, 365, 815, 535), radius=22, fill=(18, 39, 72, 255))

    # "Wave" line on monitor
    wave = [(565, 450), (600, 445), (635, 470), (670, 420), (705, 465), (740, 438), (790, 450)]
    d.line(wave, fill=(70, 232, 255, 255), width=10, joint="curve")
    for x, y in wave[::2]:
        d.ellipse((x - 8, y - 8, x + 8, y + 8), fill=(70, 232, 255, 255))

    # Stand
    d.rounded_rectangle((655, 560, 705, 630), radius=16, fill=(26, 142, 242, 255))
    d.rounded_rectangle((600, 630, 760, 680), radius=20, fill=(110, 86, 207, 255))

    img.save(path)


def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_logo(slide, x: float, y: float, w: float) -> None:
    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(x), Inches(y), width=Inches(w))


def _add_title(slide, title: str, subtitle: str | None = None) -> None:
    tx = slide.shapes.add_textbox(Inches(0.9), Inches(1.1), Inches(12.0), Inches(1.4))
    tf = tx.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(46)
    r1.font.bold = True
    r1.font.color.rgb = P.ink
    p1.space_after = Pt(6)
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(18)
        r2.font.color.rgb = P.slate


def _pill(slide, x, y, text, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.0), Inches(0.42))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    tf = shape.text_frame
    tf.text = text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].runs[0].font.size = Pt(10)
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)


def _card(slide, x, y, w, h, title: str, bullets: list[str]) -> None:
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = P.panel
    box.line.color.rgb = P.border
    box.line.width = Pt(1.0)

    tx = slide.shapes.add_textbox(Inches(x + 0.35), Inches(y + 0.28), Inches(w - 0.7), Inches(h - 0.55))
    tf = tx.text_frame
    tf.clear()
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = title
    r0.font.size = Pt(16)
    r0.font.bold = True
    r0.font.color.rgb = P.ink
    p0.space_after = Pt(8)
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(13)
        p.font.color.rgb = P.slate


def _add_footer(slide, text: str) -> None:
    tx = slide.shapes.add_textbox(Inches(0.9), Inches(7.12), Inches(12.0), Inches(0.3))
    tf = tx.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = P.muted


def build_deck() -> Path:
    _ensure_assets()
    noc_icon_path = ASSETS_DIR / "noc_engineer.png"
    if not noc_icon_path.exists():
        _make_noc_engineer_icon(noc_icon_path)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    # Slide 1: Title
    s = prs.slides.add_slide(blank)
    # Speaker notes (handy reminders when presenting)
    s.notes_slide.notes_text_frame.text = f"Demo reference: {PRESENTATION_NOTE_URL}"
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.4)
    _add_title(s, "NetCopilot", "AI-powered, evidence-first network operations copilot")
    _pill(s, 0.9, 2.35, "MULTI-VENDOR", P.blue)
    _pill(s, 3.0, 2.35, "PLAYBOOKS", P.indigo)
    _pill(s, 5.1, 2.35, "AUDIT-READY", P.blue)
    _card(
        s,
        0.9,
        3.2,
        5.9,
        3.6,
        "What it does",
        [
            "Turns chat questions into verified diagnostics",
            "Runs safe, read-only checks across devices + tools",
            "Produces clear evidence, diagnosis, and next steps",
            "Optionally creates/updates tickets with proof",
        ],
    )
    _card(
        s,
        7.05,
        3.2,
        5.35,
        3.6,
        "Who it’s for",
        [
            "NOC / operations teams",
            "Network engineering escalation",
            "Service provider & enterprise networks",
            "Multi-vendor environments",
        ],
    )
    _add_footer(s, "NetCopilot • Pitch deck")

    # Slide 2: Problem
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "The problem", "Troubleshooting is slow, inconsistent, and hard to audit")
    _card(
        s,
        0.9,
        2.5,
        6.2,
        4.6,
        "Today’s reality",
        [
            "Engineers jump between CLI, monitoring, syslog, inventory, tickets",
            "Triage steps vary by person; knowledge is tribal",
            "Incomplete context leads to false positives or slow escalation",
            "Tickets often lack evidence and clear next actions",
        ],
    )
    _card(
        s,
        7.35,
        2.5,
        5.05,
        4.6,
        "Business impact",
        [
            "Higher MTTR and longer customer impact windows",
            "More escalations due to inconsistent triage",
            "Operational load increases during incident storms",
            "Harder to onboard new engineers quickly",
        ],
    )
    _add_footer(s, "NetCopilot • Problem statement")

    # Slide 3: Solution
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "The solution", "Standardized triage using playbooks + real evidence")
    _card(
        s,
        0.9,
        2.55,
        3.95,
        4.55,
        "Ask",
        [
            "Engineer asks in plain English",
            "Slack / UI / API entry points",
            "No need to remember commands",
        ],
    )
    _card(
        s,
        5.1,
        2.55,
        3.95,
        4.55,
        "Verify",
        [
            "Runs playbook steps",
            "Calls devices + tools safely",
            "Collects raw outputs + key fields",
        ],
    )
    _card(
        s,
        9.3,
        2.55,
        3.05,
        4.55,
        "Act",
        [
            "Evidence + diagnosis",
            "Clear next steps",
            "Ticket updates (optional)",
        ],
    )
    _add_footer(s, "NetCopilot • Evidence-first workflow")

    # Slide 4: NOC Engineer (requested)
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "How it helps a NOC engineer", "From alert → diagnosis → ticket, in one workflow")
    s.shapes.add_picture(str(noc_icon_path), Inches(0.9), Inches(2.1), width=Inches(4.0))
    _card(
        s,
        5.1,
        2.2,
        7.25,
        4.85,
        "Typical flow",
        [
            "Alert arrives (Prometheus / syslog / NMS) → Slack channel",
            "Engineer runs: /netcopilot investigate",
            "NetCopilot executes link_down playbook (local + remote checks)",
            "Returns: status, evidence, likely cause, next steps",
            "Optional: creates/updates Jira with the evidence attached",
        ],
    )
    _add_footer(s, "NetCopilot • NOC experience")

    # Slide 5: Example flow (Interface Down)
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "End-to-end flow (example: Link Down)", "Playbook-driven troubleshooting with evidence")

    # Flow boxes
    x0, y0 = 0.9, 2.15
    w, h = 3.15, 0.95
    gap = 0.35
    steps = [
        ("User request", "“Investigate link down on pe1 ge-0/0/3”"),
        ("Playbook", "link_down (vendor-aware steps)"),
        ("Device checks", "Interface status/errors + LLDP neighbor"),
        ("Correlate", "Prometheus metrics + syslog search"),
        ("Answer", "Diagnosis + next steps + optional ticket"),
    ]
    for i, (t, b) in enumerate(steps):
        bx = x0 + i * (w + gap)
        box = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y0), Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = P.panel
        box.line.color.rgb = P.border
        tx = s.shapes.add_textbox(Inches(bx + 0.18), Inches(y0 + 0.12), Inches(w - 0.36), Inches(h - 0.2))
        tf = tx.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        r0 = p0.add_run()
        r0.text = t
        r0.font.bold = True
        r0.font.size = Pt(13)
        r0.font.color.rgb = P.ink
        p1 = tf.add_paragraph()
        p1.text = b
        p1.font.size = Pt(11)
        p1.font.color.rgb = P.slate
        if i < len(steps) - 1:
            arrow = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(bx + w + 0.05), Inches(y0 + 0.28), Inches(0.25), Inches(0.4))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = P.blue
            arrow.line.color.rgb = P.blue

    _card(
        s,
        0.9,
        3.7,
        11.5,
        3.35,
        "Output format (what the engineer gets)",
        [
            "Status: local admin/oper, last flap, alarms/defects, errors/drops/CRC",
            "Evidence: relevant CLI snippets + correlated telemetry/logs",
            "Likely cause: physical vs remote shutdown vs protocol issue",
            "Next steps: concrete checks or escalation actions",
        ],
    )
    _add_footer(s, "NetCopilot • Link down workflow")

    # Slide 6: Architecture
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "Architecture (high level)", "FastAPI + LLM tool orchestration + adapters")

    # Left: channels
    _card(s, 0.9, 2.2, 3.2, 1.65, "Channels", ["Web UI", "Slack / Teams", "API clients"])
    # Middle: API + orchestration
    _card(s, 4.35, 2.2, 4.2, 1.65, "FastAPI service", ["Auth0 RBAC + audit logs", "POST /chat", "Admin settings/playbooks"])
    _card(s, 4.35, 4.1, 4.2, 1.65, "LLM orchestrator", ["System prompt + tools", "Playbook-first triage", "Final narrative response"])
    # Right: adapters
    _card(
        s,
        8.85,
        2.2,
        3.55,
        3.55,
        "Adapters",
        [
            "Devices: IOS XR / Junos / Nokia",
            "Monitoring: Prometheus",
            "Logs: Syslog/Loki",
            "Inventory: YAML/NetBox/Nautobot",
            "Ticketing: Jira (+ ServiceNow)",
        ],
    )
    # Stores
    _card(s, 0.9, 4.1, 3.2, 1.65, "Local stores", ["settings.db", "playbooks.db", "alerts.db", "memory.db"])
    _add_footer(s, "NetCopilot • Architecture")

    # Slide 7: Integrations + adapter model
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "Integrations via adapters", "Add customer systems without rewriting the core")
    _card(
        s,
        0.9,
        2.55,
        6.1,
        4.55,
        "Why adapters",
        [
            "Consistent API for each capability (devices, tickets, inventory)",
            "Vendor/tool specifics live behind adapter boundaries",
            "Easier testing and faster onboarding for new integrations",
            "Customer-specific connectors become configuration, not rewrites",
        ],
    )
    _card(
        s,
        7.25,
        2.55,
        5.15,
        4.55,
        "Current integrations",
        [
            "Devices: SSH/CLI (Netmiko) + NETCONF (extensible)",
            "Monitoring: Prometheus queries",
            "Logs: Syslog/Loki query adapter",
            "Inventory: devices.yaml + Nautobot/NetBox",
            "Ticketing: Jira create/comment/close",
        ],
    )
    _add_footer(s, "NetCopilot • Adapter strategy")

    # Slide 8: Security & governance
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "Security & governance", "Read-only by default, auditable, policy-driven")
    _card(
        s,
        0.9,
        2.55,
        6.1,
        4.55,
        "Controls",
        [
            "Auth0 roles (admin vs noc) for settings + actions",
            "Safe execution: tool allowlist/denylist + CLI prefix allowlist",
            "Audit logs for requests, playbook changes, and ticket actions",
            "Secrets stored as settings; redaction supported in logs",
        ],
    )
    _card(
        s,
        7.25,
        2.55,
        5.15,
        4.55,
        "Deployment posture",
        [
            "Run in customer environment (VM/K8s) with their network access",
            "LLM provider choice (OpenAI / Bedrock / local) via LiteLLM",
            "Data minimization: send only what’s required to diagnose",
            "Future: approvals + guardrails for config changes",
        ],
    )
    _add_footer(s, "NetCopilot • Security posture")

    # Slide 9: Pilot plan
    s = prs.slides.add_slide(blank)
    _set_bg(s, P.bg)
    _add_logo(s, 0.9, 0.7, 1.0)
    _add_title(s, "Pilot plan", "Fast path to a measurable production pilot")
    _card(
        s,
        0.9,
        2.55,
        6.1,
        4.55,
        "2–4 week pilot",
        [
            "Pick 3–5 high-volume incident types (link down, BGP down, high CPU, route missing)",
            "Connect inventory + 1 alert source + 1 ticketing system",
            "Define playbooks (default + customer tweaks) and validate outputs",
            "Roll out to NOC for real incidents with guardrails",
        ],
    )
    _card(
        s,
        7.25,
        2.55,
        5.15,
        4.55,
        "Success metrics",
        [
            "MTTR reduction (median & P90)",
            "Ticket quality: evidence + next steps completeness",
            "Engineer time saved per incident",
            "Decrease in escalations for known incident classes",
        ],
    )
    _add_footer(s, "NetCopilot • Pilot + success metrics")

    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_deck()
    print(f"Created: {path}")
